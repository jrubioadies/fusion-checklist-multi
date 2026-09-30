#!/usr/bin/env python3
"""
Fusion Checklist (multi-cliente) — Comparador de configuración entre instancias
Oracle Fusion vía consultas SQL de BI Publisher.

Cada cliente tiene su perfil (instancias a comparar + usuarios) creado desde el
asistente de la propia web; las contraseñas van al llavero del sistema.

Usage:  python3 app.py            # serves http://127.0.0.1:8900
"""
import json, sys, re, os, shutil, webbrowser, io
from urllib.parse import urlparse, parse_qs
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from fusion_client import bip as fbip
import checklist as ckl
import profiles

HOST, PORT = "127.0.0.1", 8900

# Cliente abierto en esta sesión: perfil + contraseñas en memoria por entorno.
SESSION = {"client": None, "profile": None, "conns": {}}


def cfg():
    """Vista de configuración del cliente activo con la forma que espera el resto
    del código: {"environments": {nombre: {base, user, pass}}, "sql_report", ...}."""
    prof = SESSION.get("profile") or {}
    envs = {}
    for e in prof.get("environments") or []:
        conn = SESSION["conns"].get(e["name"]) or {}
        envs[e["name"]] = {"base": e.get("base", ""), "user": e.get("user", ""), "pass": conn.get("pass", "")}
    c = {"environments": envs, "datasources": prof.get("datasources") or {}}
    if prof.get("sql_report"):
        c["sql_report"] = prof["sql_report"]
    return c


def compare_envs():
    """Entornos marcados para comparar, en el orden del perfil."""
    prof = SESSION.get("profile") or {}
    return [e["name"] for e in prof.get("environments") or [] if e.get("compare", True)]


def open_client(slug, passwords=None):
    """Activa un cliente: carga su perfil y las contraseñas (las recibidas o las del
    llavero). Devuelve la lista de entornos a comparar que siguen sin contraseña."""
    prof = profiles.load(slug)
    if not prof:
        raise RuntimeError(f"No existe el cliente '{slug}'.")
    passwords = passwords or {}
    SESSION.update(client=profiles.slugify(slug), profile=prof, conns={})
    for e in prof.get("environments") or []:
        pw = passwords.get(e["name"]) or profiles.get_password(slug, e["name"])
        if pw:
            SESSION["conns"][e["name"]] = {"base": e["base"].rstrip("/"), "user": e.get("user", ""), "pass": pw}
    return [n for n in compare_envs() if n not in SESSION["conns"]]


def creds(c, env=None):
    name = env or (compare_envs() or [""])[0]
    e = (c.get("environments") or {}).get(name) or {}
    return (e.get("base", "").rstrip("/"), e.get("user", ""), e.get("pass", ""))


def report_for(c, ds):
    ds_map = c.get("datasources", {})
    return (ds_map.get(ds) if ds else None) or ds_map.get(c.get("default_ds", "")) \
        or c.get("sql_report", fbip.SQL_REPORT_DEFAULT)


def dict_query(c, ds, sql, env=None):
    base, user, pw = creds(c, env)
    data = fbip.run_sql(base, user, pw, report_for(c, ds), sql)
    return fbip.parse_rows(data)


BU_CANDIDATES = ["BUSINESS_UNIT_ID", "BU_ID", "ORG_ID", "PRC_BU_ID",
                 "REQUISITIONING_BU_ID", "MANUFACTURING_BU_ID"]


def san_obj(s):
    return re.sub(r"[^A-Za-z0-9_%$#./]", "", (s or ""))[:60].upper()


def bu_column_for(c, ds, sql, env):
    tables = re.findall(r"\b(?:from|join)\s+([a-zA-Z0-9_$]+)", sql or "", re.I)
    tables = list(dict.fromkeys(san_obj(t) for t in tables if san_obj(t)))[:6]
    if not tables:
        return None
    inlist = ",".join(f"'{t}'" for t in tables)
    candlist = ",".join(f"'{x}'" for x in BU_CANDIDATES)
    q = (f"SELECT DISTINCT column_name FROM all_tab_columns "
         f"WHERE table_name IN ({inlist}) AND column_name IN ({candlist})")
    try:
        _, rows = dict_query(c, ds, q, env)
    except Exception:
        return None
    found = {r.get("COLUMN_NAME", "") for r in rows}
    for x in BU_CANDIDATES:
        if x in found:
            return x
    return None


def prepare_sql(sql, user, bu_id, limit, bu_col=None):
    sql = (sql or "").strip().rstrip(";")
    sql = re.sub(r":xdo_user_name\b", "'" + (user or "").replace("'", "''") + "'", sql, flags=re.I)
    safe_bu = re.sub(r"[^0-9A-Za-z_-]", "", str(bu_id)) if bu_id else ""
    if bu_id:
        sql = re.sub(r":bu_id\b", safe_bu, sql, flags=re.I)
    if bu_id and bu_col and re.match(r"(?is)\s*(with|select)\b", sql):
        # La proyección (COUNT(*) o lista de columnas) puede no exponer la columna
        # de BU. Envolvemos la tabla origen con SELECT * para que la columna esté
        # disponible en el filtro, conservando la proyección original.
        m = re.match(r"(?is)\s*select\s+(.*?)\s+from\b(.*)$", sql)
        if m:
            proj, tail = m.group(1).strip(), m.group(2).strip()
            # Extraer cualquier ROWNUM<=N y aplicarlo DESPUÉS del filtro de BU;
            # si no, limita las filas antes de filtrar y puede dar 0 falsos.
            mn = re.search(r"(?is)rownum\s*<=\s*(\d+)", tail)
            rn = mn.group(1) if mn else None
            if rn:
                tail = re.sub(r"(?is)\s+and\s+rownum\s*<=\s*\d+", "", tail)
                tail = re.sub(r"(?is)\brownum\s*<=\s*\d+\s+and\s+", "", tail)
                tail = re.sub(r"(?is)\s+where\s+rownum\s*<=\s*\d+\b", "", tail)
                tail = tail.strip()
            sql = (f"SELECT {proj} FROM (SELECT * FROM {tail}) fbip_bu "
                   f"WHERE fbip_bu.{bu_col} = {safe_bu}")
            if rn:
                sql += f" AND ROWNUM <= {rn}"
        else:
            sql = f"SELECT * FROM (\n{sql}\n) fbip_bu WHERE fbip_bu.{bu_col} = {safe_bu}"
    if limit and str(limit).lower() != "all":
        n = int(re.sub(r"[^0-9]", "", str(limit)) or "0")
        if n > 0 and re.match(r"(?is)\s*(with|select)\b", sql):
            sql = f"SELECT * FROM (\n{sql}\n) WHERE ROWNUM <= {n}"
    return sql


def esc_html(s):
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


# ── FBDI: extracción de delta de configuración entre entornos ───────────

def _norm(s):
    return (s or "").strip().upper()


def _key_banks(r):
    return (_norm(r.get("COUNTRY")), _norm(r.get("BANK_NAME")))


def _key_branches(r):
    return _norm(r.get("BIC")) or (_norm(r.get("BANK_NAME")), _norm(r.get("BRANCH_NUMBER")))


def _key_accounts(r):
    return _norm(r.get("IBAN")) or _norm(r.get("ACCOUNT_NUMBER"))


# Cada item FBDI define su plantilla oficial y las "hojas" a extraer.
FBDI_SPECS = {
    "bank_accounts": {
        "steps": [26, 28, 29, 101, 102, 103],
        "title": "Bancos, Sucursales y Cuentas (Cash Management)",
        "template_task": "Create Banks, Branches, and Accounts in Spreadsheet",
        "upload_task": "Upload Banks, Branches, and Accounts",
        "note": "Revisa Legal Entity y usos (Payable/Receivable) antes de subir; el número de cuenta/IBAN se extrae en claro.",
        "sheets": [
            {"name": "Banks", "keyfn": _key_banks,
             "sql": "SELECT COUNTRY_NAME COUNTRY, BANK_NAME, BANK_NUMBER, BANK_NAME_ALT ALT_NAME, "
                    "SHORT_BANK_NAME SHORT_NAME, DESCRIPTION FROM CE_BANKS_V"},
            {"name": "Branches", "keyfn": _key_branches,
             "sql": "SELECT BANK_NAME, BANK_BRANCH_NAME BRANCH_NAME, BRANCH_NUMBER, "
                    "BANK_BRANCH_NAME_ALT ALT_NAME, EFT_SWIFT_CODE BIC, BANK_BRANCH_TYPE BRANCH_TYPE, "
                    "COUNTRY_NAME COUNTRY FROM CE_BANK_BRANCHES_V"},
            {"name": "Accounts", "keyfn": _key_accounts,
             "sql": "SELECT bk.BANK_NAME, br.BANK_BRANCH_NAME BRANCH_NAME, ba.BANK_ACCOUNT_NAME ACCOUNT_NAME, "
                    "ba.BANK_ACCOUNT_NUM ACCOUNT_NUMBER, ba.IBAN_NUMBER IBAN, ba.CURRENCY_CODE CURRENCY, "
                    "ba.BANK_ACCOUNT_TYPE ACCOUNT_TYPE, ba.AP_USE_ALLOWED_FLAG PAYABLE, "
                    "ba.AR_USE_ALLOWED_FLAG RECEIVABLE FROM CE_BANK_ACCOUNTS ba, CE_BANKS_V bk, CE_BANK_BRANCHES_V br "
                    "WHERE ba.BANK_ID=bk.BANK_PARTY_ID AND ba.BANK_BRANCH_ID=br.BRANCH_PARTY_ID"},
        ],
    },
    "document_sequences": {
        "steps": [19, 60],
        "title": "Document Sequences (Payables / Receivables)",
        "template_task": "REST (documentSequences)",
        "upload_task": "REST POST documentSequences",
        "note": "Migración por REST. El ModuleId (GUID de módulo) se reutiliza del origen. Sin DELETE por REST (rollback por UI).",
        "sheets": [],  # solo REST (sin vía Excel/BI)
    },
}


def _make_keyfn(key_cols):
    def kf(r):
        return tuple(_norm(r.get(k)) for k in key_cols)
    return kf


# Items FBDI de una sola tabla: delta + Excel (extracción validada contra la instancia).
FBDI_SIMPLE = [
    {"step": 2,  "title": "Legal Entities",              "sheet": "LegalEntities",  "key": ["LEGAL_ENTITY_IDENTIFIER"],
     "sql": "SELECT NAME, LEGAL_ENTITY_IDENTIFIER, TRANSACTING_ENTITY_FLAG FROM XLE_ENTITY_PROFILES"},
    {"step": 8,  "title": "Business Units",              "sheet": "BusinessUnits",  "key": ["BU_NAME"],
     "sql": "SELECT BU_NAME, STATUS FROM FUN_ALL_BUSINESS_UNITS_V WHERE status='A'"},
    {"step": 1,  "title": "Business Units (Procurement)", "sheet": "BusinessUnits", "key": ["BU_NAME"],
     "sql": "SELECT BU_NAME, STATUS FROM FUN_ALL_BUSINESS_UNITS_V WHERE status='A'"},
    {"step": 13, "title": "Payment Terms (AP)",          "sheet": "PaymentTerms",   "key": ["NAME"],
     "sql": "SELECT NAME, ENABLED_FLAG, START_DATE_ACTIVE FROM AP_TERMS"},
    {"step": 21, "title": "Payment Methods",             "sheet": "PaymentMethods", "key": ["PAYMENT_METHOD_CODE"],
     "sql": "SELECT PAYMENT_METHOD_CODE, PAYMENT_METHOD_NAME FROM IBY_PAYMENT_METHODS_VL"},
    {"step": 40, "title": "Tax Rates",                   "sheet": "TaxRates",       "key": ["TAX_RATE_CODE", "PERCENTAGE_RATE"],
     "sql": "SELECT TAX_RATE_CODE, PERCENTAGE_RATE, ACTIVE_FLAG FROM ZX_RATES_B"},
    {"step": 42, "title": "Tax Regimes",                 "sheet": "TaxRegimes",     "key": ["TAX_REGIME_CODE"],
     "sql": "SELECT TAX_REGIME_CODE, EFFECTIVE_FROM FROM ZX_REGIMES_B"},
    {"step": 49, "title": "Receivables Activities",      "sheet": "RecvActivities", "key": ["NAME"],
     "sql": "SELECT NAME, TYPE, STATUS FROM AR_RECEIVABLES_TRX_ALL"},
    {"step": 52, "title": "Transaction Types",           "sheet": "TrxTypes",       "key": ["NAME"],
     "sql": "SELECT NAME, TYPE, STATUS FROM RA_CUST_TRX_TYPES_ALL"},
    {"step": 55, "title": "Payment Terms (AR)",          "sheet": "PaymentTermsAR", "key": ["NAME"],
     "sql": "SELECT NAME, START_DATE_ACTIVE, END_DATE_ACTIVE FROM RA_TERMS"},
    {"step": 62, "title": "Standard Memo Lines",         "sheet": "MemoLines",      "key": ["NAME"],
     "sql": "SELECT NAME, LINE_TYPE FROM AR_MEMO_LINES_ALL_B"},
    {"step": 63, "title": "Receipt Classes and Methods", "sheet": "ReceiptMethods", "key": ["NAME"],
     "sql": "SELECT NAME, RECEIPT_CLASS_ID FROM AR_RECEIPT_METHODS"},
    {"step": 65, "title": "Receipt Sources",             "sheet": "ReceiptSources", "key": ["NAME"],
     "sql": "SELECT NAME, TYPE, START_DATE_ACTIVE FROM AR_BATCH_SOURCES_ALL"},
    {"step": 66, "title": "Transaction Sources",         "sheet": "TrxSources",     "key": ["NAME"],
     "sql": "SELECT NAME, STATUS, DESCRIPTION FROM RA_BATCH_SOURCES_ALL"},
    {"step": 85, "title": "Asset Books",                 "sheet": "AssetBooks",     "key": ["BOOK_TYPE_CODE"],
     "sql": "SELECT BOOK_TYPE_CODE, BOOK_TYPE_NAME, BOOK_CLASS FROM FA_BOOK_CONTROLS"},
    {"step": 97, "title": "Asset Categories",            "sheet": "AssetCategories", "key": ["SEGMENT1", "SEGMENT2", "SEGMENT3"],
     "sql": "SELECT SEGMENT1, SEGMENT2, SEGMENT3, ENABLED_FLAG FROM FA_CATEGORIES_B"},
    {"step": 5,  "title": "Inventory Organizations",     "sheet": "InvOrgs",        "key": ["ORGANIZATION_CODE"],
     "sql": "SELECT ORGANIZATION_CODE, ORGANIZATION_NAME FROM INV_ORGANIZATION_DEFINITIONS_V"},
    {"step": 9,  "title": "Locations",                   "sheet": "Locations",      "key": ["INTERNAL_LOCATION_CODE"],
     "sql": "SELECT INTERNAL_LOCATION_CODE, COUNTRY FROM PER_LOCATIONS"},
]
FBDI_SIMPLE_BY_STEP = {it["step"]: it for it in FBDI_SIMPLE}


def _simple_spec(it):
    return {
        "steps": [it["step"]],
        "title": it["title"],
        "template_task": "FBDI / Configuration Workbook (según módulo)",
        "upload_task": "Import / Load correspondiente",
        "note": "Delta para preparar la carga. Confirma la plantilla oficial del módulo antes de subir.",
        "sheets": [{"name": it["sheet"], "keyfn": _make_keyfn(it["key"]), "sql": it["sql"]}],
    }


def fbdi_spec_for_step(step):
    for spec in FBDI_SPECS.values():
        if step in spec["steps"]:
            return spec
    it = FBDI_SIMPLE_BY_STEP.get(step)
    if it:
        return _simple_spec(it)
    return None


def fbdi_compute_delta(c, spec, source, target):
    """Para cada hoja: extrae origen y destino, y devuelve solo lo que falta en destino."""
    ds = "FSCM (Financials/SCM)"
    out = []
    for sh in spec["sheets"]:
        _, src = dict_query(c, ds, sh["sql"], source)
        _, tgt = dict_query(c, ds, sh["sql"], target)
        tkeys = set(sh["keyfn"](r) for r in tgt)
        delta = [r for r in src if sh["keyfn"](r) not in tkeys]
        cols = list(delta[0].keys()) if delta else (list(src[0].keys()) if src else [])
        out.append({"name": sh["name"], "columns": cols, "rows": delta,
                    "src_count": len(src), "tgt_count": len(tgt), "missing": len(delta)})
    return out


def fbdi_build_xlsx(sheets):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        ws = wb.create_sheet((sh["name"] or "Sheet")[:31])
        ws.append(sh["columns"])
        for r in sh["rows"]:
            ws.append([r.get(col, "") for col in sh["columns"]])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ── Carga por REST API (crear en destino lo que falta) ──────────────────
# Los recursos cashBanks/cashBankBranches/cashBankAccounts se referencian por
# claves naturales (nombre/número), así que se crean en cadena sin IDs internos.
# La creación funciona; el borrado por REST NO es fiable → DRY RUN obligatorio.

REST_API_PATH = "/fscmRestApi/resources/11.13.18.05/"

try:  # las instancias dev usan verify=False; evitar ruido de InsecureRequestWarning
    fbip.requests.packages.urllib3.disable_warnings()
except Exception:
    pass

# Objetos migrables por REST para el spec bank_accounts (orden = dependencia).
FBDI_REST_OBJECTS = {
    "bank_accounts": [
        {"name": "Banks", "resource": "cashBanks",
         "fields": "BankName,BankNameAlt,BankNumber,CountryName,Description",
         "keyfn": lambda r: (_norm(r.get("CountryName")), _norm(r.get("BankName"))),
         "keep": ["BankName", "BankNameAlt", "BankNumber", "CountryName", "Description"]},
        {"name": "Branches", "resource": "cashBankBranches",
         "fields": "BankName,BankNumber,BankBranchName,BranchNumber,BankBranchNameAlt,EFTSWIFTCode,CountryName",
         "keyfn": lambda r: _norm(r.get("EFTSWIFTCode")) or (_norm(r.get("BankName")), _norm(r.get("BranchNumber"))),
         "keep": ["BankName", "BankNumber", "BankBranchName", "BranchNumber", "BankBranchNameAlt", "EFTSWIFTCode", "CountryName"]},
        {"name": "Accounts", "resource": "cashBankAccounts",
         "fields": "BankName,BankBranchName,BankAccountName,BankAccountNumber,IBANNumber,CurrencyCode,"
                   "LegalEntityName,AccountType,ApUseAllowedFlag,ArUseAllowedFlag",
         "keyfn": lambda r: _norm(r.get("IBANNumber")) or _norm(r.get("BankAccountNumber")),
         "keep": ["BankName", "BankBranchName", "BankAccountName", "BankAccountNumber", "IBANNumber",
                  "CurrencyCode", "LegalEntityName", "AccountType", "ApUseAllowedFlag", "ArUseAllowedFlag"]},
    ],
    "document_sequences": [
        {"name": "DocumentSequences", "resource": "documentSequences",
         "fields": "Name,Type,DeterminantType,ApplicationId,ApplicationShortName,ModuleId,StartDate,EndDate,InitialValue",
         "keyfn": lambda r: _norm(r.get("Name")),
         "keep": ["Name", "Type", "DeterminantType", "ApplicationId", "ApplicationShortName",
                  "ModuleId", "StartDate", "EndDate", "InitialValue"]},
    ],
}


def fbdi_rest_objects(spec):
    for key, s in FBDI_SPECS.items():
        if s is spec:
            return FBDI_REST_OBJECTS.get(key, [])
    return []


def _rest_get_all(base, user, pw, resource, fields):
    items, offset = [], 0
    while True:
        url = (base.rstrip("/") + REST_API_PATH + resource +
               f"?limit=500&offset={offset}" + (f"&fields={fields}" if fields else ""))
        r = fbip.requests.get(url, auth=(user, pw), headers={"REST-Framework-Version": "1"},
                              verify=False, timeout=120)
        r.raise_for_status()
        j = r.json()
        items += j.get("items", [])
        if not j.get("hasMore"):
            break
        offset += 500
    return items


def _rest_payload(row, keep):
    return {k: row.get(k) for k in keep if row.get(k) not in (None, "")}


def _rest_err(resp):
    try:
        j = resp.json()
        return j.get("detail") or j.get("title") or (resp.text[:300])
    except Exception:
        return resp.text[:300]


def fbdi_rest_plan(c, spec, source, target):
    """DRY RUN: extrae por REST origen y destino y devuelve los payloads a crear."""
    objs = fbdi_rest_objects(spec)
    sb, su, sp = creds(c, source)
    tb, tu, tp = creds(c, target)
    out = []
    for obj in objs:
        src = _rest_get_all(sb, su, sp, obj["resource"], obj["fields"])
        tgt = _rest_get_all(tb, tu, tp, obj["resource"], obj["fields"])
        tkeys = set(obj["keyfn"](r) for r in tgt)
        payloads = [_rest_payload(r, obj["keep"]) for r in src if obj["keyfn"](r) not in tkeys]
        out.append({"name": obj["name"], "resource": obj["resource"],
                    "src_count": len(src), "tgt_count": len(tgt),
                    "missing": len(payloads), "payloads": payloads})
    return out


def fbdi_rest_execute(c, spec, source, target):
    """Crea en destino los registros que faltan. Devuelve un log por registro."""
    plan = fbdi_rest_plan(c, spec, source, target)
    tb, tu, tp = creds(c, target)
    H = {"REST-Framework-Version": "1", "Content-Type": "application/json"}
    log, summary = [], []
    for obj in plan:
        ok_n, err_n = 0, 0
        url = tb.rstrip("/") + REST_API_PATH + obj["resource"]
        for pl in obj["payloads"]:
            label = pl.get("BankAccountName") or pl.get("BankBranchName") or pl.get("BankName") or "?"
            try:
                r = fbip.requests.post(url, auth=(tu, tp), headers=H, data=json.dumps(pl),
                                       verify=False, timeout=60)
                if r.status_code in (200, 201):
                    ok_n += 1
                    log.append({"object": obj["name"], "key": label, "status": r.status_code, "ok": True})
                else:
                    err_n += 1
                    log.append({"object": obj["name"], "key": label, "status": r.status_code,
                                "ok": False, "detail": _rest_err(r)})
            except Exception as e:
                err_n += 1
                log.append({"object": obj["name"], "key": label, "status": 0, "ok": False, "detail": str(e)[:200]})
        summary.append({"name": obj["name"], "created": ok_n, "failed": err_n, "total": len(obj["payloads"])})
    return {"summary": summary, "log": log}


# ── HTML page ──────────────────────────────────────────────────────────

PAGE = r"""<!doctype html><html lang="es"><head><meta charset="utf-8">
<title>Fusion Checklist</title>
<meta name="viewport" content="width=device-width,initial-scale=1">
<style>
  :root{--bg:#0b0f14;--panel:#0d1117;--border:#22272e;--fg:#e6edf3;--mut:#7d8590;--acc:#1f6feb;--ok:#3fb950;--err:#f85149;--warn:#d29922}
  *{box-sizing:border-box} body{margin:0;background:var(--bg);color:var(--fg);font:13px system-ui,Segoe UI,Roboto,sans-serif}
  header{position:sticky;top:0;z-index:10;background:var(--panel);border-bottom:1px solid var(--border);padding:10px 16px;display:flex;gap:14px;align-items:center;flex-wrap:wrap}
  header h1{font-size:16px;margin:0;font-weight:800} header .sub{color:var(--mut);font-size:12px}
  .actions{display:flex;gap:8px;margin-left:auto;align-items:center}
  .actions button{background:var(--acc);color:#fff;border:0;border-radius:7px;padding:8px 16px;font-weight:700;cursor:pointer;font-size:13px}
  .actions button:disabled{opacity:.5}
  .actions #exportXlsx{background:#1a4b2e;color:#3fb950;border:1px solid #2ea043}
  .actions .progress{font-size:11px;color:var(--mut)}
  .logout{background:#161b22!important;border:1px solid var(--border)!important;color:var(--mut)!important;font-size:11px!important;padding:6px 10px!important}
  .wrap{padding:16px}
  table.ck{border-collapse:collapse;width:100%;font:12px ui-monospace,Menlo,monospace}
  table.ck th{position:sticky;top:49px;background:#11161d;border-bottom:2px solid var(--border);padding:8px 10px;text-align:left;color:#adbac7;font-size:11px;text-transform:uppercase;letter-spacing:.4px}
  table.ck td{border-bottom:1px solid var(--border);padding:7px 10px;white-space:nowrap}
  table.ck tr:hover{background:#1c2430}
  table.ck .mod{color:#a371f7;font-weight:700;font-size:11px}
  table.ck .task{color:var(--fg);white-space:normal;max-width:300px}
  table.ck .cnt{text-align:right;font-weight:700;min-width:60px;cursor:pointer}
  table.ck .cnt:hover{text-decoration:underline;color:var(--acc)}
  table.ck .cnt.loading{color:var(--mut);font-style:italic;font-weight:400}
  table.ck .cnt.err{color:var(--err);font-size:10px;font-weight:400;cursor:help}
  .match{color:var(--ok)} .mismatch{color:var(--warn)} .zero{color:var(--mut)}
  .badge-method{font-size:10px;font-weight:700;padding:2px 6px;border-radius:4px;white-space:nowrap}
  .badge-method.rest{background:#1a4b2e;color:#3fb950} .badge-method.fbdi{background:#2d1f0e;color:#d29922}
  .badge-method.rest-fbdi{background:#0c2d4a;color:#58a6ff} .badge-method.ui{background:#1c1c1c;color:#7d8590}
  .diff{font-size:10px;color:var(--warn);margin-left:4px}
  .fbdi-btn{background:#0c2d4a;border:1px solid #1f6feb;color:#58a6ff;border-radius:5px;padding:1px 5px;font-size:11px;cursor:pointer;margin-left:5px}
  .fbdi-btn:hover{background:#123a5e}
  #fbdiModal .detail-hd select{background:#0b0f14;color:var(--fg);border:1px solid var(--border);border-radius:6px;padding:4px 7px;font-size:12px}
  #fbdiModal .detail-hd button{margin-left:0}
  #fbdiModal .detail-hd #fbdiCalc{background:var(--acc);color:#fff;border:0}
  #fbdiModal .detail-hd #fbdiDl{background:#1a4b2e;color:#3fb950;border:1px solid #2ea043}
  #fbdiModal .detail-hd #fbdiDl:disabled{opacity:.5}
  #fbdiModal .detail-hd #fbdiPlan{background:#0c2d4a;color:#58a6ff;border:1px solid #1f6feb}
  #fbdiModal .detail-hd #fbdiExec{background:#5a1e1e;color:#ff7b72;border:1px solid #da3633}
  #fbdiModal .detail-hd #fbdiExec:disabled{opacity:.45}
  .detail-modal{display:none;position:fixed;inset:0;background:rgba(0,0,0,.6);z-index:200;padding:24px}
  .detail-modal.open{display:block}
  .detail-box{background:var(--panel);border:1px solid var(--border);border-radius:12px;height:100%;display:flex;flex-direction:column;overflow:hidden}
  .detail-hd{display:flex;gap:10px;align-items:center;padding:10px 14px;border-bottom:1px solid var(--border);background:#11161d;flex-wrap:wrap}
  .detail-hd b{font-size:14px} .detail-hd .env{background:#6e40c9;color:#fff;padding:2px 8px;border-radius:6px;font-size:11px;font-weight:700}
  .detail-hd button{background:#161b22;color:var(--fg);border:1px solid var(--border);border-radius:7px;padding:6px 10px;cursor:pointer;margin-left:auto}
  .detail-body{flex:1;overflow:auto;padding:0}
  .detail-body table{border-collapse:collapse;width:100%;font:11.5px ui-monospace,Menlo,monospace}
  .detail-body th{position:sticky;top:0;background:#11161d;border-bottom:1px solid var(--border);padding:6px 8px;text-align:left;color:#adbac7}
  .detail-body td{border-bottom:1px solid var(--border);padding:5px 8px;white-space:nowrap;max-width:300px;overflow:hidden;text-overflow:ellipsis}
  .detail-body tr:nth-child(even){background:#0e131a}
  .filter-row{padding:8px 16px;display:flex;gap:10px;align-items:center;border-bottom:1px solid var(--border)}
  .filter-row label{font-size:11px;color:var(--mut)}
  .filter-row select,.filter-row input{background:#0b0f14;color:var(--fg);border:1px solid var(--border);border-radius:6px;padding:5px 8px;font-size:12px}
  .loginov{display:none;position:fixed;inset:0;background:radial-gradient(1200px 600px at 50% -10%,#16213a,#0b0f14);z-index:300;align-items:flex-start;justify-content:center;overflow:auto;padding:5vh 16px}
  .loginov.open{display:flex}
  .wz{width:860px;max-width:100%;background:var(--panel);border:1px solid var(--border);border-radius:14px;padding:24px 26px 20px;box-shadow:0 24px 70px rgba(0,0,0,.6)}
  .wz h2{font-size:20px;margin:0;font-weight:800} .wz .lgsub{color:var(--mut);font-size:12.5px;margin:3px 0 16px}
  .wz .steps{display:flex;gap:6px;margin:0 0 16px;font-size:11px;color:var(--mut)}
  .wz .steps span{padding:3px 9px;border:1px solid var(--border);border-radius:99px} .wz .steps span.on{border-color:#6e40c9;color:var(--fg);background:#231a3a}
  .wz label.f{display:block;font-size:11px;text-transform:uppercase;letter-spacing:.5px;color:var(--mut);margin:10px 0 4px}
  .wz input[type=text],.wz input[type=password],.wz select{background:#0b0f14;color:var(--fg);border:1px solid var(--border);border-radius:8px;padding:8px 10px;outline:0;font-size:13px;width:100%}
  .wz input:focus,.wz select:focus{border-color:#6e40c9}
  .wz table{width:100%;border-collapse:collapse;margin-top:6px} .wz td,.wz th{padding:5px 4px;vertical-align:top;text-align:left}
  .wz th{font-size:10.5px;color:var(--mut);text-transform:uppercase;letter-spacing:.4px;font-weight:600}
  .wz .st{font-size:11.5px;line-height:1.45;color:var(--mut);max-width:240px} .wz .st.ok{color:var(--ok)} .wz .st.err{color:var(--err)} .wz .st.warn{color:var(--warn)}
  .wz button{background:#161b22;color:var(--fg);border:1px solid var(--border);border-radius:8px;padding:7px 12px;cursor:pointer;font-size:12.5px;white-space:nowrap}
  .wz button.pri{background:#6e40c9;border-color:#6e40c9;color:#fff;font-weight:700}
  .wz button.dan{color:var(--err)} .wz button:disabled{opacity:.55;cursor:default}
  .wz .row{display:flex;gap:8px;align-items:center;flex-wrap:wrap} .wz .foot{display:flex;gap:8px;margin-top:18px;align-items:center}
  .wz .foot .sp{flex:1} .wz .lgerr{color:var(--err);font-size:12.5px;margin-top:10px;min-height:16px}
  .wz .chk{display:flex;gap:7px;align-items:center;font-size:12px;margin-top:10px}
  .wz .cl{display:flex;gap:10px;align-items:center;border:1px solid var(--border);border-radius:10px;padding:10px 12px;margin-bottom:8px}
  .wz .cl b{font-size:14px} .wz .cl .envs{color:var(--mut);font-size:11.5px;flex:1}
  .wz .hint{color:var(--mut);font-size:11px;margin-top:14px;line-height:1.5;border-top:1px solid var(--border);padding-top:10px}
  .wz details{margin-top:12px;font-size:12px;color:var(--mut)}
  @media (max-width:640px){ .wz table,.wz tbody,.wz tr,.wz td{display:block} .wz thead{display:none} .wz tr{border:1px solid var(--border);border-radius:8px;padding:6px;margin-bottom:8px} }
  /* ── Loader estilo Minecraft (Creeper bailando) ── */
  @keyframes mcdance{0%,100%{transform:rotate(-9deg) translateY(0)}50%{transform:rotate(9deg) translateY(-7px)}}
  .proc{display:flex;flex-direction:column;align-items:center;justify-content:center;gap:10px;padding:26px 20px;text-align:center}
  .mc{display:inline-block;position:relative;width:48px;height:48px;background:#66a83f;box-shadow:inset 0 0 0 2px #4c7d2e;image-rendering:pixelated;animation:mcdance .5s ease-in-out infinite}
  .mc::before{content:"";position:absolute;left:0;top:0;width:6px;height:6px;background:transparent;box-shadow:6px 6px #17300f,12px 6px #17300f,30px 6px #17300f,36px 6px #17300f,6px 12px #17300f,12px 12px #17300f,30px 12px #17300f,36px 12px #17300f,18px 18px #17300f,24px 18px #17300f,12px 24px #17300f,18px 24px #17300f,24px 24px #17300f,30px 24px #17300f,12px 30px #17300f,18px 30px #17300f,24px 30px #17300f,30px 30px #17300f,12px 36px #17300f,30px 36px #17300f}
  .proc .msg{font-weight:800;color:var(--fg);font-size:14px}
  .proc .sub{font-size:11px;color:var(--mut)}
  .proc.mini{flex-direction:row;gap:8px;padding:0;display:inline-flex;vertical-align:middle}
  .proc.mini .mc{width:22px;height:22px;animation-duration:.45s}
  .proc.mini .mc::before{box-shadow:none}
  .proc.mini .msg{font-size:12px;font-weight:700}
</style></head><body>
<div id="loginOv" class="loginov"><div class="wz" id="wz"></div></div>
<header>
  <h1>📋 Fusion Checklist</h1>
  <span class="sub" id="clientName">Comparador de entornos Oracle Fusion</span>
  <span class="sub" id="lgWho" style="color:var(--ok)"></span>
  <div class="actions">
    <span class="progress" id="progress"></span>
    <button id="runAll">▶ Ejecutar comparación</button>
    <button id="exportXlsx" title="Excel con conteos y entidades comparadas entre entornos (respeta los filtros)">⬇ Exportar Excel</button>
    <button id="editClient" class="logout" title="Instancias y usuarios del cliente">⚙ instancias</button>
    <button id="logout" class="logout">cambiar cliente</button>
  </div>
</header>
<div class="filter-row">
  <label>Módulo <select id="fMod"><option value="">Todos</option></select></label>
  <label>Método <select id="fMethod"><option value="">Todos</option><option value="REST">REST</option><option value="FBDI">FBDI</option><option value="REST+FBDI">REST+FBDI</option><option value="UI">UI</option></select></label>
  <label>Business Unit <select id="fBU"><option value="">(todas)</option></select></label>
  <label>Buscar <input id="fSearch" placeholder="filtrar tarea…" style="min-width:200px"/></label>
  <button id="loadBU" style="background:#161b22;color:var(--fg);border:1px solid var(--border);border-radius:6px;padding:5px 10px;cursor:pointer;font-size:11px">↻ Cargar BUs</button>
  <span id="buStatus" style="font-size:11px;color:var(--mut)"></span>
</div>
<div class="wrap"><table class="ck" id="tbl"><thead><tr>
  <th>#</th><th>Módulo</th><th>Tarea</th><th>Responsable</th><th>Método</th>
  <th id="envHeads" hidden></th><th>Estado</th>
</tr></thead><tbody id="tbody"></tbody></table></div>
<div class="detail-modal" id="detailModal">
  <div class="detail-box">
    <div class="detail-hd"><b id="detTtl">Detalle</b><span class="env" id="detEnv"></span><span id="detInfo" style="color:var(--mut);font-size:11px"></span><button id="detClose">✕ cerrar</button></div>
    <div class="detail-body" id="detBody"></div>
  </div>
</div>
<div class="detail-modal" id="fbdiModal">
  <div class="detail-box">
    <div class="detail-hd">
      <b id="fbdiTtl">FBDI</b>
      <span id="fbdiInfo" style="color:var(--mut);font-size:11px"></span>
      <label style="font-size:11px;color:var(--mut);margin-left:auto">Origen <select id="fbdiSrc"></select></label>
      <label style="font-size:11px;color:var(--mut)">Destino <select id="fbdiTgt"></select></label>
      <button id="fbdiCalc">Calcular delta</button>
      <button id="fbdiDl" disabled>⬇ Descargar Excel</button>
      <button id="fbdiPlan">🧪 Simular REST</button>
      <button id="fbdiExec" disabled>⚡ Cargar REST</button>
      <button id="fbdiClose">✕ cerrar</button>
    </div>
    <div class="detail-body" id="fbdiBody"></div>
  </div>
</div>
<script>
function showJsErr(msg){
  var b=document.getElementById('jserr');
  if(!b){b=document.createElement('div');b.id='jserr';b.style.cssText='position:fixed;bottom:0;left:0;right:0;background:#5a1e1e;color:#ff7b72;font:12px ui-monospace,monospace;padding:6px 10px;z-index:99999;white-space:pre-wrap;border-top:2px solid #da3633';document.body.appendChild(b);}
  b.textContent='⚠ '+msg;
}
window.addEventListener('error',function(e){ showJsErr('JS error: '+e.message+'  ('+(e.filename||'').split('/').pop()+':'+e.lineno+')'); });
window.addEventListener('unhandledrejection',function(e){ showJsErr('Promesa rechazada: '+((e.reason&&(e.reason.stack||e.reason.message))||e.reason)); });
let ENVS=[], AUTH={};
let ITEMS=[], results={};
const $=s=>document.querySelector(s);
function getSelectedBU(){ return $('#fBU').value||''; }

// ── asistente de cliente: 1) instancias 2) usuarios y contraseñas ──
const escH=s=>String(s==null?'':s).replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
const post=(url,body)=>fetch(url,{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify(body||{})}).then(r=>r.json());
let WZ=null;  // {slug, name, environments:[{name,base,user,compare,has_password,pass,probe,test}], sql_report, fbdi_source, fbdi_target}

function wzOpen(){ $('#loginOv').classList.add('open'); }
function wzSteps(n){ return '<div class="steps">'+['Cliente','1 · Instancias','2 · Usuarios y contraseñas'].map((t,i)=>'<span class="'+(i===n?'on':'')+'">'+t+'</span>').join('')+'</div>'; }

function wzClients(){
  const cs=AUTH.clients||[];
  $('#wz').innerHTML='<h2>🔐 Fusion Checklist</h2><div class="lgsub">Elige el cliente o da de alta uno nuevo con sus instancias Oracle Fusion.</div>'+wzSteps(0)+
    (cs.length?cs.map(c=>'<div class="cl"><b>'+escH(c.name)+'</b><span class="envs">'+escH((c.environments||[]).join(' · '))+'</span>'+
      '<button class="pri" data-open="'+escH(c.slug)+'">Abrir</button><button data-edit="'+escH(c.slug)+'">Editar</button><button class="dan" data-del="'+escH(c.slug)+'">✕</button></div>').join('')
      :'<div class="st">Todavía no hay clientes configurados.</div>')+
    '<div class="foot"><button class="pri" id="wzNew">+ Nuevo cliente</button><span class="sp"></span></div><div class="lgerr" id="wzErr"></div>'+
    '<div class="hint">Los perfiles se guardan en <code>~/.config/fusion-checklist/clients/</code>; las contraseñas '+(AUTH.keyring?'en el llavero del sistema.':'<b>solo en memoria</b> (instala <code>keyring</code> para recordarlas).')+'</div>';
  $('#wzNew').onclick=()=>{ WZ={slug:'',name:'',sql_report:'',environments:[newEnv('PROD'),newEnv('TEST')]}; wzInstances(); };
  $('#wz').querySelectorAll('[data-open]').forEach(b=>b.onclick=()=>openClient(b.dataset.open));
  $('#wz').querySelectorAll('[data-edit]').forEach(b=>b.onclick=()=>editClient(b.dataset.edit,1));
  $('#wz').querySelectorAll('[data-del]').forEach(b=>b.onclick=async()=>{
    if(b.dataset.confirm!=='1'){ b.dataset.confirm='1'; b.textContent='¿Eliminar?'; return; }
    await post('/api/clients/delete',{slug:b.dataset.del}); await refreshAuth(); wzClients();
  });
  wzOpen();
}
function newEnv(name){ return {name:name||'',base:'',user:'',compare:true,pass:'',has_password:false}; }

async function editClient(slug,step){
  const d=await fetch('/api/clients/get?slug='+encodeURIComponent(slug)).then(r=>r.json());
  if(d.error){ $('#wzErr').textContent='⚠ '+d.error; return; }
  WZ=d; WZ.environments.forEach(e=>{ e.pass=''; if(e.compare===undefined) e.compare=true; });
  step===2?wzCreds():wzInstances(); wzOpen();
}
async function openClient(slug){
  const d=await post('/api/clients/open',{slug});
  if(d.error){ $('#wzErr').textContent='⚠ '+d.error; return; }
  if(d.missing&&d.missing.length) return editClient(slug,2);  // faltan contraseñas → paso 2
  enterApp();
}

// Paso 1: instancias que se van a comparar
function wzInstances(){
  const E=WZ.environments;
  const rows=E.map((e,i)=>{
    const p=e.probe, st=!p?'<span class="st">sin evaluar</span>':p.loading?'<span class="st">evaluando…</span>':
      '<span class="st '+(p.ok?'ok':'err')+'">'+(p.ok?'✓ ':'⚠ ')+escH(p.message)+(p.latency_ms!=null?' · '+p.latency_ms+' ms':'')+'</span>';
    return '<tr><td style="width:110px"><input type="text" data-k="name" data-i="'+i+'" value="'+escH(e.name)+'" placeholder="PROD"/></td>'+
      '<td><input type="text" data-k="base" data-i="'+i+'" value="'+escH(e.base)+'" placeholder="https://xxxx.fa.em2.oraclecloud.com" spellcheck="false"/></td>'+
      '<td style="text-align:center"><input type="checkbox" data-k="compare" data-i="'+i+'" '+(e.compare?'checked':'')+' title="Incluir en la comparación"/></td>'+
      '<td><button data-probe="'+i+'">Evaluar</button></td><td>'+st+'</td>'+
      '<td><button class="dan" data-rm="'+i+'" title="Quitar">✕</button></td></tr>';
  }).join('');
  const envOpts=sel=>E.filter(e=>e.name).map(e=>'<option'+(e.name===sel?' selected':'')+'>'+escH(e.name)+'</option>').join('');
  $('#wz').innerHTML='<h2>'+(WZ.slug?'Editar cliente':'Nuevo cliente')+'</h2><div class="lgsub">Indica las instancias del cliente y cuáles quieres comparar. «Evaluar» comprueba que la URL es una instancia Fusion accesible (sin credenciales).</div>'+wzSteps(1)+
    '<label class="f">Nombre del cliente</label><input type="text" id="wzName" value="'+escH(WZ.name)+'" placeholder="p. ej. Acme S.A." '+(WZ.slug?'disabled':'')+'/>'+
    '<label class="f">Instancias</label><table><thead><tr><th>Nombre</th><th>URL</th><th>Comparar</th><th></th><th>Estado</th><th></th></tr></thead><tbody>'+rows+'</tbody></table>'+
    '<div class="row" style="margin-top:8px"><button id="wzAdd">+ Añadir instancia</button><button id="wzProbeAll">Evaluar todas</button></div>'+
    '<details><summary>Opciones avanzadas</summary>'+
      '<label class="f">Report ejecutor SQL (BI Publisher)</label><input type="text" id="wzReport" value="'+escH(WZ.sql_report||'')+'" placeholder="/Custom/SQLTools/SQLConReport.xdo"/>'+
      '<div class="row" style="margin-top:6px"><label class="f" style="margin:0">Origen FBDI por defecto <select id="wzSrc">'+envOpts(WZ.fbdi_source||(E[0]||{}).name)+'</select></label>'+
      '<label class="f" style="margin:0">Destino <select id="wzTgt">'+envOpts(WZ.fbdi_target||(E[E.length-1]||{}).name)+'</select></label></div></details>'+
    '<div class="foot"><button id="wzBack">← Clientes</button><span class="sp"></span><button class="pri" id="wzNext">Siguiente: usuarios →</button></div><div class="lgerr" id="wzErr"></div>';
  const W=$('#wz');
  W.querySelectorAll('input[data-k]').forEach(inp=>inp.oninput=inp.onchange=()=>{
    const e=E[+inp.dataset.i], k=inp.dataset.k;
    if(k==='compare') e.compare=inp.checked; else { e[k]=k==='name'?inp.value.toUpperCase():inp.value; if(k==='base') e.probe=null; }
    if(k==='name') inp.value=e.name;
  });
  $('#wzName').oninput=ev=>WZ.name=ev.target.value;
  $('#wzReport').oninput=ev=>WZ.sql_report=ev.target.value.trim();
  $('#wzSrc').onchange=ev=>WZ.fbdi_source=ev.target.value; $('#wzTgt').onchange=ev=>WZ.fbdi_target=ev.target.value;
  $('#wzAdd').onclick=()=>{ E.push(newEnv('')); wzInstances(); };
  W.querySelectorAll('[data-rm]').forEach(b=>b.onclick=()=>{ E.splice(+b.dataset.rm,1); wzInstances(); });
  W.querySelectorAll('[data-probe]').forEach(b=>b.onclick=()=>probeEnv(+b.dataset.probe));
  $('#wzProbeAll').onclick=()=>E.forEach((_,i)=>probeEnv(i));
  $('#wzBack').onclick=async()=>{ await refreshAuth(); wzClients(); };
  $('#wzNext').onclick=()=>{
    const err=validateInstances(); if(err){ $('#wzErr').textContent='⚠ '+err; return; }
    wzCreds();
  };
}
async function probeEnv(i){
  const e=WZ.environments[i]; if(!e.base.trim()){ e.probe={ok:false,message:'Falta la URL'}; return wzInstances(); }
  e.probe={loading:true}; wzInstances();
  try{
    const d=await post('/api/setup/probe',{base:e.base});
    e.probe=d; if(d.base) e.base=d.base;
    if(!e.name && d.suggested_name && !WZ.environments.some(x=>x.name===d.suggested_name)) e.name=d.suggested_name;
  }catch(ex){ e.probe={ok:false,message:String(ex)}; }
  wzInstances();
}
function validateInstances(){
  const E=WZ.environments;
  if(!(WZ.name||'').trim()) return 'Pon un nombre al cliente.';
  if(E.some(e=>!e.name.trim()||!e.base.trim())) return 'Cada instancia necesita nombre y URL.';
  if(new Set(E.map(e=>e.name)).size!==E.length) return 'Los nombres de instancia deben ser únicos.';
  if(E.filter(e=>e.compare).length<2) return 'Marca al menos dos instancias para comparar.';
  return '';
}

// Paso 2: usuario y contraseña de cada instancia
function wzCreds(){
  const E=WZ.environments;
  const rows=E.map((e,i)=>{
    const t=e.test; let st='<span class="st">'+(e.has_password?'contraseña guardada':'sin probar')+'</span>';
    if(t&&t.loading) st='<span class="st">probando…</span>';
    else if(t&&!t.ok) st='<span class="st err">⚠ '+escH(t.message)+'</span>';
    else if(t&&t.ok){
      const q=t.sql||{};
      st='<span class="st ok">✓ login correcto</span><br>'+(q.ok?'<span class="st ok">✓ SQL: '+escH(q.ledgers)+' ledgers · '+escH(q.bus)+' BUs · '+escH(q.legal_entities)+' LEs</span>'
        :'<span class="st warn">⚠ SQL: '+escH(q.message||'no disponible')+'</span>');
    }
    return '<tr><td style="width:120px"><b>'+escH(e.name)+'</b>'+(e.compare?'':' <span class="st">(no compara)</span>')+'<div class="st">'+escH(e.base.replace(/^https?:\/\//,''))+'</div></td>'+
      '<td><input type="text" data-k="user" data-i="'+i+'" value="'+escH(e.user)+'" placeholder="usuario" autocomplete="off" spellcheck="false"/></td>'+
      '<td><input type="password" data-k="pass" data-i="'+i+'" value="'+escH(e.pass)+'" placeholder="'+(e.has_password?'•••••• (sin cambios)':'contraseña')+'" autocomplete="new-password"/></td>'+
      '<td><button data-test="'+i+'">Probar</button></td><td>'+st+'</td></tr>';
  }).join('');
  $('#wz').innerHTML='<h2>'+escH(WZ.name)+'</h2><div class="lgsub">Usuario y contraseña de Fusion para cada instancia. «Probar» valida el login y que el report ejecutor de SQL está desplegado.</div>'+wzSteps(2)+
    '<table><thead><tr><th>Instancia</th><th>Usuario</th><th>Contraseña</th><th></th><th>Resultado</th></tr></thead><tbody>'+rows+'</tbody></table>'+
    '<div class="row" style="margin-top:8px"><button id="wzSame">Usar el primer usuario/contraseña en todas</button><button id="wzTestAll">Probar todas</button></div>'+
    '<label class="chk"><input type="checkbox" id="wzRemember" '+(AUTH.keyring?'checked':'disabled')+'/> Guardar contraseñas en el llavero del sistema'+(AUTH.keyring?'':' (no disponible)')+'</label>'+
    '<div class="foot"><button id="wzBack">← Instancias</button><span class="sp"></span><button class="pri" id="wzSave">Guardar y abrir</button></div><div class="lgerr" id="wzErr"></div>'+
    '<div class="hint">Las credenciales solo se usan en este equipo para conectar con Fusion; no se envían a ningún servidor externo.</div>';
  const W=$('#wz');
  W.querySelectorAll('input[data-k]').forEach(inp=>inp.oninput=()=>{ const e=E[+inp.dataset.i]; e[inp.dataset.k]=inp.value; e.test=null; });
  W.querySelectorAll('[data-test]').forEach(b=>b.onclick=()=>testEnv(+b.dataset.test));
  $('#wzTestAll').onclick=()=>E.forEach((_,i)=>testEnv(i));
  $('#wzSame').onclick=()=>{ const f=E[0]; E.forEach(e=>{ e.user=f.user; if(f.pass) e.pass=f.pass; e.test=null; }); wzCreds(); };
  $('#wzBack').onclick=wzInstances;
  $('#wzSave').onclick=saveClient;
}
async function testEnv(i){
  const e=WZ.environments[i];
  if(!e.user.trim()||(!e.pass&&!e.has_password)){ e.test={ok:false,message:'Falta usuario o contraseña'}; return wzCreds(); }
  e.test={loading:true}; wzCreds();
  try{ e.test=await post('/api/setup/test',{base:e.base,user:e.user,pass:e.pass,slug:WZ.slug,env:e.name,sql_report:WZ.sql_report}); }
  catch(ex){ e.test={ok:false,message:String(ex)}; }
  wzCreds();
}
async function saveClient(){
  const E=WZ.environments, btn=$('#wzSave');
  const miss=E.filter(e=>e.compare&&(!e.user.trim()||(!e.pass&&!e.has_password))).map(e=>e.name);
  if(miss.length){ $('#wzErr').textContent='⚠ Falta usuario o contraseña en: '+miss.join(', '); return; }
  btn.disabled=true; btn.textContent='Guardando…';
  const profile={slug:WZ.slug,name:WZ.name.trim(),sql_report:WZ.sql_report||undefined,
    fbdi_source:WZ.fbdi_source,fbdi_target:WZ.fbdi_target,
    environments:E.map(e=>({name:e.name,base:e.base,user:e.user.trim(),compare:!!e.compare}))};
  const passwords={}; E.forEach(e=>{ if(e.pass) passwords[e.name]=e.pass; });
  try{
    const d=await post('/api/clients/save',{profile,passwords,remember:$('#wzRemember').checked});
    if(!d.ok){ $('#wzErr').textContent='⚠ '+(d.error||'Error'); return; }
    if(d.missing&&d.missing.length){ $('#wzErr').textContent='⚠ Sin contraseña para: '+d.missing.join(', '); return; }
    location.reload();
  }catch(ex){ $('#wzErr').textContent='⚠ '+ex; }
  finally{ btn.disabled=false; btn.textContent='Guardar y abrir'; }
}

async function refreshAuth(){ AUTH=await fetch('/api/auth').then(r=>r.json()); return AUTH; }
function enterApp(){ location.reload(); }
function renderEnvHeads(){
  const h=$('#envHeads'); h.insertAdjacentHTML('beforebegin',ENVS.map(e=>'<th>'+escH(e)+'</th>').join('')); h.remove();
}
async function checkAuth(){
  try{ await refreshAuth(); }catch(e){ AUTH={clients:[]}; }
  if(AUTH.logged_in){
    ENVS=AUTH.envs; renderEnvHeads();
    $('#loginOv').classList.remove('open');
    $('#clientName').textContent='🏢 '+AUTH.client.name+' · '+ENVS.join(' / ');
    $('#lgWho').textContent='🔒 '+(AUTH.users||[]).join(', ');
    loadItems();
  } else if(AUTH.client){ editClient(AUTH.client.slug,2); }
  else wzClients();
}
$('#logout').onclick=async()=>{ await fetch('/api/logout',{method:'POST'}); location.reload(); };
$('#editClient').onclick=()=>AUTH.client&&editClient(AUTH.client.slug,1);

// ── checklist ──
async function loadBUs(){
  const btn=$('#loadBU'), sel=$('#fBU'), st=$('#buStatus');
  btn.disabled=true; btn.textContent='cargando…'; if(st) st.textContent='⏳ cargando…';
  try{
    const r=await fetch('/api/checklist/bus',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({env:ENVS[0]})});
    const d=await r.json();
    sel.innerHTML='<option value="">(todas)</option>';
    if(d.error){ if(st) st.textContent='⚠ '+d.error; }
    (d.bus||[]).forEach(b=>{const o=document.createElement('option');o.value=b.id;o.textContent=b.name+' ('+b.id+')';sel.appendChild(o);});
    if(st && !d.error) st.textContent='✓ '+(d.bus||[]).length+' BUs';
  }catch(e){
    if(st) st.textContent='⚠ '+e;
  }
  btn.disabled=false; btn.textContent='↻ Cargar BUs';
}

async function loadItems(){
  const d=await fetch('/api/checklist/items').then(r=>r.json());
  ITEMS=d.items||[];
  const mods=[...new Set(ITEMS.map(i=>i.module))];
  mods.forEach(m=>{const o=document.createElement('option');o.value=m;o.textContent=m;$('#fMod').appendChild(o);});
  try{ renderTable(); }catch(err){ showJsErr('renderTable: '+(err.stack||err.message||err)); }
  loadBUs();
}

function filteredItems(){
  const mod=$('#fMod').value, q=$('#fSearch').value.toLowerCase(), mf=$('#fMethod').value;
  return ITEMS.filter(it=>{
    if(mod && it.module!==mod) return false;
    if(mf && (it.config_method||'UI')!==mf && !(mf==='REST'&&(it.config_method||'').includes('REST')) && !(mf==='FBDI'&&(it.config_method||'').includes('FBDI'))) return false;
    if(q && !it.task.toLowerCase().includes(q)) return false;
    return true;
  });
}

function renderTable(){
  const filtered=filteredItems();
  const tbody=$('#tbody'); tbody.innerHTML='';
  filtered.forEach(it=>{
    const tr=document.createElement('tr');
    const envCells=ENVS.map(env=>{
      const key=it.step+'_'+env;
      const r=results[key];
      if(!r) return '<td class="cnt" data-step="'+it.step+'" data-env="'+env+'">—</td>';
      if(r.error) return '<td class="cnt err" data-step="'+it.step+'" data-env="'+env+'" title="'+esc(r.error)+'">⚠</td>';
      return '<td class="cnt" data-step="'+it.step+'" data-env="'+env+'">'+r.count+'</td>';
    }).join('');
    const counts=ENVS.map(env=>results[it.step+'_'+env]).filter(r=>r&&!r.error).map(r=>parseInt(r.count)||0);
    let status='';
    if(counts.length===ENVS.length){
      const allSame=counts.every(c=>c===counts[0]);
      if(allSame && counts[0]===0) status='<span class="zero">vacío</span>';
      else if(allSame) status='<span class="match">✓ iguales</span>';
      else status='<span class="mismatch" style="cursor:pointer" data-diffstep="'+it.step+'">⚠ diferencias</span>';
    }
    const meth=it.config_method||'';
    const mCls=meth==='REST'?'rest':meth==='FBDI'?'fbdi':meth==='REST+FBDI'?'rest-fbdi':'ui';
    const methBadge='<span class="badge-method '+mCls+'">'+esc(meth||'UI')+'</span>'+(it.fbdi?'<button class="fbdi-btn" data-fbdi="'+it.step+'" title="Extraer delta entre entornos y descargar plantilla">📥 FBDI</button>':'');
    tr.innerHTML='<td>'+it.step+'</td><td class="mod">'+esc(it.module)+'</td><td class="task">'+esc(it.task)+'</td><td>'+esc(it.responsible||'')+'</td><td>'+methBadge+'</td>'+envCells+'<td>'+status+'</td>';
    tbody.appendChild(tr);
  });
  tbody.querySelectorAll('.cnt').forEach(td=>td.onclick=()=>{
    const step=parseInt(td.dataset.step), env=td.dataset.env;
    if(td.classList.contains('err')) return;
    showDetail(step, env);
  });
  tbody.querySelectorAll('[data-diffstep]').forEach(el=>el.onclick=()=>showDiff(parseInt(el.dataset.diffstep)));
  tbody.querySelectorAll('.fbdi-btn').forEach(b=>b.onclick=(e)=>{e.stopPropagation();openFbdi(parseInt(b.dataset.fbdi));});
}

function esc(s){return String(s??'').replace(/[&<>]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;'}[c]));}
// ── Loader estilo Minecraft (Creeper bailando) ──
const PROC_MSGS=['Picando bloques…','Generando chunks…','Crafteando la respuesta…','¡Cuidado, un Creeper! 💥','Minando datos de Oracle…','Bloque a bloque, ya casi…'];
function procMsg(){return PROC_MSGS[Math.floor(Math.random()*PROC_MSGS.length)];}
function proc(sub,mini){
  const m=procMsg();
  const mc='<span class="mc"></span>';
  if(mini) return '<span class="proc mini">'+mc+'<span class="msg">'+esc(sub||('Procesando… '+m))+'</span></span>';
  return '<div class="proc">'+mc+'<div class="msg">Procesando… ⛏️</div><div class="sub">'+esc(sub||m)+'</div></div>';
}

async function runAll(){
  const btn=$('#runAll'); btn.disabled=true; btn.textContent='⛏️ minando…';
  const bu=getSelectedBU();
  let done=0, total=ITEMS.length*ENVS.length;
  $('#progress').innerHTML='<span class="proc mini" style="display:inline-flex"><span class="mc"></span><span class="msg">Procesando <span id="progTxt">0/'+total+'</span> · '+procMsg()+'</span></span>';
  results={};
  for(const it of ITEMS){
    const promises=ENVS.map(async env=>{
      try{
        const r=await fetch('/api/checklist/run',{method:'POST',headers:{'Content-Type':'application/json'},
          body:JSON.stringify({step:it.step,env,mode:'count',bu})});
        const d=await r.json();
        results[it.step+'_'+env]=d;
      }catch(e){ results[it.step+'_'+env]={error:''+e,count:null}; }
      done++; const pt=document.getElementById('progTxt'); if(pt)pt.textContent=done+'/'+total;
      renderTable();
    });
    await Promise.all(promises);
  }
  btn.disabled=false; btn.textContent='▶ Ejecutar comparación';
  $('#progress').innerHTML='✓ completado 🟩✨';
}

async function exportXlsx(){
  const btn=$('#exportXlsx'), steps=filteredItems().map(i=>i.step), bu=getSelectedBU();
  if(!steps.length){ $('#progress').textContent='⚠ No hay tareas visibles que exportar'; return; }
  btn.disabled=true; btn.textContent='⛏️ exportando…';
  $('#progress').innerHTML=proc('Consultando '+steps.length+' tareas × '+ENVS.length+' entornos para el Excel…',true);
  try{
    const r=await fetch('/api/checklist/export',{method:'POST',headers:{'Content-Type':'application/json'},
      body:JSON.stringify({steps,bu})});
    if(!r.ok) throw new Error('HTTP '+r.status);
    if((r.headers.get('Content-Type')||'').indexOf('json')>=0){ const j=await r.json(); throw new Error(j.error||'Error'); }
    const blob=await r.blob();
    const a=document.createElement('a'); a.href=URL.createObjectURL(blob);
    a.download='Fusion_Checklist'+(bu?'_BU'+bu:'')+'.xlsx'; document.body.appendChild(a); a.click();
    a.remove(); URL.revokeObjectURL(a.href);
    $('#progress').textContent='✓ Excel generado';
  }catch(e){ $('#progress').textContent='⚠ Error al exportar: '+e.message; }
  finally{ btn.disabled=false; btn.textContent='⬇ Exportar Excel'; }
}

async function showDetail(step, env){
  const item=ITEMS.find(i=>i.step===step);
  if(!item) return;
  const bu=getSelectedBU();
  $('#detTtl').textContent=item.task;
  $('#detEnv').textContent=env;
  $('#detBody').innerHTML=proc('Cargando detalle de '+env+'…');
  $('#detailModal').classList.add('open');
  try{
    const r=await fetch('/api/checklist/run',{method:'POST',headers:{'Content-Type':'application/json'},
      body:JSON.stringify({step,env,mode:'detail',bu})});
    const d=await r.json();
    if(d.error){ $('#detBody').innerHTML='<div style="padding:20px;color:var(--err)">'+esc(d.error)+'</div>'; return; }
    $('#detInfo').textContent=(d.rows||[]).length+' registros';
    if(!d.rows||!d.rows.length){ $('#detBody').innerHTML='<div style="padding:20px;color:var(--mut)">Sin registros.</div>'; return; }
    let h='<table><thead><tr>'+d.columns.map(c=>'<th>'+esc(c)+'</th>').join('')+'</tr></thead><tbody>';
    d.rows.forEach(r=>{ h+='<tr>'+d.columns.map(c=>'<td title="'+esc(r[c])+'">'+esc(r[c])+'</td>').join('')+'</tr>'; });
    $('#detBody').innerHTML=h+'</tbody></table>';
  }catch(e){ $('#detBody').innerHTML='<div style="padding:20px;color:var(--err)">'+esc(''+e)+'</div>'; }
}

async function showDiff(step){
  const item=ITEMS.find(i=>i.step===step);
  if(!item) return;
  const bu=getSelectedBU();
  $('#detTtl').textContent='Comparación: '+item.task;
  $('#detEnv').textContent='TODOS';
  $('#detInfo').textContent='Cargando…';
  $('#detBody').innerHTML=proc('Obteniendo detalle de los '+ENVS.length+' entornos…');
  $('#detailModal').classList.add('open');
  const allData={};
  await Promise.all(ENVS.map(async env=>{
    try{
      const r=await fetch('/api/checklist/run',{method:'POST',headers:{'Content-Type':'application/json'},
        body:JSON.stringify({step,env,mode:'detail',bu})});
      allData[env]=await r.json();
    }catch(e){ allData[env]={error:''+e}; }
  }));
  const cols=ENVS.map(e=>allData[e]).filter(d=>d.columns&&d.columns.length).map(d=>d.columns)[0]||[];
  if(!cols.length){
    const anyErr=ENVS.map(e=>allData[e]).find(d=>d&&d.error);
    const msg=anyErr?('⚠ Error: '+esc(anyErr.error)):('Sin registros que comparar'+(getSelectedBU()?' para la Business Unit seleccionada.':'.'));
    $('#detBody').innerHTML='<div style="padding:20px;color:'+(anyErr?'var(--err)':'var(--mut)')+'">'+msg+'</div>'; return;
  }
  const keyCol=cols[0];
  const rowMap={};
  const allKeys=[];
  ENVS.forEach(env=>{
    const d=allData[env];
    if(!d||d.error||!d.rows) return;
    d.rows.forEach(row=>{
      const k=String(row[keyCol]||'').trim();
      if(!rowMap[k]){ rowMap[k]={}; allKeys.push(k); }
      rowMap[k][env]=row;
    });
  });
  const uniqueKeys=[...new Set(allKeys)];
  let h='<table><thead><tr><th>'+esc(keyCol)+'</th>';
  const ctxCol=cols.length>1?cols[1]:null;
  if(ctxCol) h+='<th>'+esc(ctxCol)+'</th>';
  ENVS.forEach(env=>{ h+='<th style="text-align:center">'+env+'</th>'; });
  h+='</tr></thead><tbody>';
  uniqueKeys.forEach(k=>{
    const envPresent=ENVS.map(env=>!!rowMap[k][env]);
    const allPresent=envPresent.every(Boolean);
    const rowClass=allPresent?'':'style="background:#1a1200"';
    h+='<tr '+rowClass+'><td title="'+esc(k)+'">'+esc(k)+'</td>';
    if(ctxCol){
      const sample=ENVS.map(env=>rowMap[k][env]).find(Boolean);
      h+='<td>'+esc(sample?sample[ctxCol]:'')+'</td>';
    }
    ENVS.forEach(env=>{
      if(rowMap[k][env]) h+='<td style="text-align:center;color:var(--ok)">✓</td>';
      else h+='<td style="text-align:center;color:var(--err);font-weight:700">✗</td>';
    });
    h+='</tr>';
  });
  h+='</tbody></table>';
  const presentInAll=uniqueKeys.filter(k=>ENVS.every(env=>rowMap[k][env])).length;
  const onlyInSome=uniqueKeys.length-presentInAll;
  $('#detInfo').textContent=uniqueKeys.length+' registros únicos · '+presentInAll+' en todos · '+onlyInSome+' con diferencias';
  $('#detBody').innerHTML=h;
}

$('#detClose').onclick=()=>$('#detailModal').classList.remove('open');
$('#detailModal').addEventListener('click',e=>{if(e.target.id==='detailModal')$('#detailModal').classList.remove('open');});
document.addEventListener('keydown',e=>{if(e.key==='Escape')$('#detailModal').classList.remove('open');});
$('#runAll').onclick=runAll;
$('#exportXlsx').onclick=exportXlsx;
$('#fMod').onchange=renderTable;
$('#fMethod').onchange=renderTable;
$('#fSearch').oninput=renderTable;
$('#loadBU').onclick=loadBUs;

// ── FBDI: extraer delta entre entornos y descargar plantilla ──
let fbdiStep=null;
function fillEnvSelect(sel, def){ sel.innerHTML=''; ENVS.forEach(e=>{const o=document.createElement('option');o.value=e;o.textContent=e;if(e===def)o.selected=true;sel.appendChild(o);}); }
function openFbdi(step){
  fbdiStep=step;
  const it=ITEMS.find(i=>i.step===step);
  $('#fbdiTtl').textContent='📥 FBDI · '+(it?it.task:('paso '+step));
  fillEnvSelect($('#fbdiSrc'),AUTH.fbdi_source||ENVS[0]); fillEnvSelect($('#fbdiTgt'),AUTH.fbdi_target||ENVS[ENVS.length-1]);
  $('#fbdiInfo').textContent='';
  const canRest=!!(it&&it.fbdi_rest), canBI=!!(it&&it.fbdi_bi);
  $('#fbdiCalc').style.display=canBI?'':'none';
  $('#fbdiDl').style.display=canBI?'':'none';
  $('#fbdiPlan').style.display=canRest?'':'none';
  $('#fbdiExec').style.display=canRest?'':'none';
  $('#fbdiDl').disabled=true; $('#fbdiExec').disabled=true; fbdiPlanCache=null;
  const hint=[canBI?'<b>Calcular delta</b> (ver/descargar Excel)':'', canRest?'<b>🧪 Simular REST</b> (dry-run de carga por API)':''].filter(Boolean).join(' o ');
  $('#fbdiBody').innerHTML='<div style="padding:20px;color:var(--mut)">Elige <b>Origen</b> (config correcta) y <b>Destino</b> (al que falta). Luego: '+hint+'.</div>';
  $('#fbdiModal').classList.add('open');
}
async function fbdiCalc(){
  const src=$('#fbdiSrc').value, tgt=$('#fbdiTgt').value;
  if(src===tgt){ $('#fbdiBody').innerHTML='<div style="padding:20px;color:var(--warn)">Origen y destino deben ser distintos.</div>'; return; }
  $('#fbdiBody').innerHTML=proc('Calculando delta '+src+' → '+tgt+'…');
  const btn=$('#fbdiCalc'); btn.disabled=true; btn.textContent='calculando…'; $('#fbdiDl').disabled=true;
  try{
    const d=await fetch('/api/fbdi/delta',{method:'POST',headers:{'Content-Type':'application/json'},
      body:JSON.stringify({step:fbdiStep,source:src,target:tgt})}).then(r=>r.json());
    if(d.error){ $('#fbdiBody').innerHTML='<div style="padding:20px;color:var(--err)">'+esc(d.error)+'</div>'; return; }
    renderFbdi(d);
    $('#fbdiDl').disabled=!(d.sheets||[]).some(s=>s.missing>0);
  }catch(e){ $('#fbdiBody').innerHTML='<div style="padding:20px;color:var(--err)">'+esc(''+e)+'</div>'; }
  finally{ btn.disabled=false; btn.textContent='Calcular delta'; }
}
function renderFbdi(d){
  let h='<div style="padding:10px 14px;border-bottom:1px solid var(--border);font-size:12px;line-height:1.6">'
    +'<b>'+esc(d.title)+'</b><br>'
    +'<span style="color:var(--mut)">Plantilla: <b>'+esc(d.template_task)+'</b> · Subida: <b>'+esc(d.upload_task)+'</b></span>'
    +(d.note?'<br><span style="color:var(--warn)">⚠ '+esc(d.note)+'</span>':'')+'</div>';
  (d.sheets||[]).forEach(s=>{
    h+='<div style="padding:10px 14px 4px"><b>'+esc(s.name)+'</b> — faltan <b style="color:var(--warn)">'+s.missing+'</b> en '+esc(d.target)
      +' <span style="color:var(--mut)">('+esc(d.source)+': '+s.src_count+' · '+esc(d.target)+': '+s.tgt_count+')</span></div>';
    if(s.missing>0){
      h+='<div style="overflow:auto;max-height:280px;padding:0 14px 12px"><table><thead><tr>'
        +s.columns.map(c=>'<th>'+esc(c)+'</th>').join('')+'</tr></thead><tbody>';
      s.rows.slice(0,1000).forEach(r=>{ h+='<tr>'+s.columns.map(c=>'<td title="'+esc(r[c])+'">'+esc(r[c])+'</td>').join('')+'</tr>'; });
      h+='</tbody></table>'+(s.missing>1000?'<div style="color:var(--mut);padding:4px">… mostrando 1000 de '+s.missing+' (el Excel los incluye todos)</div>':'')+'</div>';
    }
  });
  const total=(d.sheets||[]).reduce((a,s)=>a+s.missing,0);
  $('#fbdiInfo').textContent=total+' registros a crear en '+d.target;
  $('#fbdiBody').innerHTML=h;
}
async function fbdiDownload(){
  const src=$('#fbdiSrc').value, tgt=$('#fbdiTgt').value;
  const btn=$('#fbdiDl'); const old=btn.textContent; btn.disabled=true; btn.textContent='generando…';
  try{
    const r=await fetch('/api/fbdi/download',{method:'POST',headers:{'Content-Type':'application/json'},
      body:JSON.stringify({step:fbdiStep,source:src,target:tgt})});
    if(!r.ok) throw new Error('HTTP '+r.status);
    const ct=r.headers.get('Content-Type')||'';
    if(ct.indexOf('json')>=0){ const j=await r.json(); throw new Error(j.error||'Error'); }
    const blob=await r.blob();
    const a=document.createElement('a'); a.href=URL.createObjectURL(blob);
    a.download='FBDI_'+src+'_to_'+tgt+'.xlsx'; document.body.appendChild(a); a.click();
    a.remove(); URL.revokeObjectURL(a.href);
  }catch(e){ alert('Error al descargar: '+e); }
  finally{ btn.disabled=false; btn.textContent=old; }
}
// ── Carga por REST (dry-run + ejecución) ──
let fbdiPlanCache=null;
async function fbdiPlan(){
  const src=$('#fbdiSrc').value, tgt=$('#fbdiTgt').value;
  if(src===tgt){ $('#fbdiBody').innerHTML='<div style="padding:20px;color:var(--warn)">Origen y destino deben ser distintos.</div>'; return; }
  $('#fbdiBody').innerHTML=proc('🧪 Simulando REST '+src+' → '+tgt+' (sin escribir nada)…');
  const btn=$('#fbdiPlan'); btn.disabled=true; btn.textContent='simulando…'; $('#fbdiExec').disabled=true; fbdiPlanCache=null;
  try{
    const d=await fetch('/api/fbdi/rest_plan',{method:'POST',headers:{'Content-Type':'application/json'},
      body:JSON.stringify({step:fbdiStep,source:src,target:tgt})}).then(r=>r.json());
    if(d.error){ $('#fbdiBody').innerHTML='<div style="padding:20px;color:var(--err)">'+esc(d.error)+'</div>'; return; }
    fbdiPlanCache={source:src,target:tgt};
    renderPlan(d);
    const total=(d.plan||[]).reduce((a,o)=>a+o.missing,0);
    $('#fbdiExec').disabled=(total===0);
  }catch(e){ $('#fbdiBody').innerHTML='<div style="padding:20px;color:var(--err)">'+esc(''+e)+'</div>'; }
  finally{ btn.disabled=false; btn.textContent='🧪 Simular REST'; }
}
function renderPlan(d){
  const total=(d.plan||[]).reduce((a,o)=>a+o.missing,0);
  let h='<div style="padding:10px 14px;border-bottom:1px solid var(--border);font-size:12px">'
    +'<b>🧪 DRY RUN · '+esc(d.source)+' → '+esc(d.target)+'</b> — se crearían <b style="color:#ff7b72">'+total+'</b> registros por REST. '
    +'<span style="color:var(--mut)">No se ha escrito nada.</span></div>';
  (d.plan||[]).forEach(o=>{
    h+='<div style="padding:10px 14px 4px"><b>'+esc(o.name)+'</b> ('+esc(o.resource)+') — a crear <b style="color:#ff7b72">'+o.missing+'</b>'
      +' <span style="color:var(--mut)">('+esc(d.source)+': '+o.src_count+' · '+esc(d.target)+': '+o.tgt_count+')</span></div>';
    if(o.missing>0){
      const cols=Object.keys(o.payloads[0]||{});
      h+='<div style="overflow:auto;max-height:240px;padding:0 14px 12px"><table><thead><tr>'+cols.map(c=>'<th>'+esc(c)+'</th>').join('')+'</tr></thead><tbody>';
      o.payloads.slice(0,500).forEach(p=>{ h+='<tr>'+cols.map(c=>'<td title="'+esc(p[c])+'">'+esc(p[c])+'</td>').join('')+'</tr>'; });
      h+='</tbody></table>'+(o.missing>500?'<div style="color:var(--mut);padding:4px">… mostrando 500 de '+o.missing+'</div>':'')+'</div>';
    }
  });
  $('#fbdiInfo').textContent=total+' a crear (dry-run)';
  $('#fbdiBody').innerHTML=h;
}
async function fbdiExec(){
  if(!fbdiPlanCache){ alert('Ejecuta primero «Simular REST».'); return; }
  const {source:src,target:tgt}=fbdiPlanCache;
  if(src!==$('#fbdiSrc').value||tgt!==$('#fbdiTgt').value){ alert('Cambió el origen/destino. Vuelve a simular.'); return; }
  if(!confirm('⚠ Vas a CREAR registros por REST en '+tgt+' (origen '+src+').\\n\\nEl borrado por REST no es fiable — revisa el dry-run antes.\\n\\n¿Continuar?')) return;
  const btn=$('#fbdiExec'); btn.disabled=true; const old=btn.textContent; btn.textContent='cargando…';
  $('#fbdiBody').innerHTML=proc('⚡ Creando registros en '+tgt+' (banks → branches → accounts)…');
  try{
    const d=await fetch('/api/fbdi/rest_execute',{method:'POST',headers:{'Content-Type':'application/json'},
      body:JSON.stringify({step:fbdiStep,source:src,target:tgt,confirm:true})}).then(r=>r.json());
    if(d.error){ $('#fbdiBody').innerHTML='<div style="padding:20px;color:var(--err)">'+esc(d.error)+'</div>'; return; }
    renderExec(d);
  }catch(e){ $('#fbdiBody').innerHTML='<div style="padding:20px;color:var(--err)">'+esc(''+e)+'</div>'; }
  finally{ btn.textContent=old; btn.disabled=false; }
}
function renderExec(d){
  let h='<div style="padding:10px 14px;border-bottom:1px solid var(--border);font-size:12px"><b>Resultado carga REST · '+esc(d.source)+' → '+esc(d.target)+'</b><br>';
  (d.summary||[]).forEach(s=>{ h+='<span style="margin-right:14px">'+esc(s.name)+': <b style="color:var(--ok)">'+s.created+'</b> creados'+(s.failed?' · <b style="color:var(--err)">'+s.failed+'</b> con error':'')+'</span>'; });
  h+='</div><div style="overflow:auto;max-height:340px;padding:6px 14px 12px"><table><thead><tr><th>Objeto</th><th>Registro</th><th>Estado</th><th>Detalle</th></tr></thead><tbody>';
  (d.log||[]).forEach(l=>{ h+='<tr><td>'+esc(l.object)+'</td><td>'+esc(l.key)+'</td>'
    +'<td style="color:'+(l.ok?'var(--ok)':'var(--err)')+'">'+(l.ok?'✓ '+l.status:'✗ '+l.status)+'</td>'
    +'<td title="'+esc(l.detail||'')+'">'+esc(l.detail||'')+'</td></tr>'; });
  h+='</tbody></table></div>';
  const created=(d.summary||[]).reduce((a,s)=>a+s.created,0), failed=(d.summary||[]).reduce((a,s)=>a+s.failed,0);
  $('#fbdiInfo').textContent=created+' creados · '+failed+' con error';
  $('#fbdiExec').disabled=true;
  $('#fbdiBody').innerHTML=h;
}
$('#fbdiClose').onclick=()=>$('#fbdiModal').classList.remove('open');
$('#fbdiModal').addEventListener('click',e=>{if(e.target.id==='fbdiModal')$('#fbdiModal').classList.remove('open');});
$('#fbdiCalc').onclick=fbdiCalc;
$('#fbdiDl').onclick=fbdiDownload;
$('#fbdiPlan').onclick=fbdiPlan;
$('#fbdiExec').onclick=fbdiExec;

checkAuth();
</script></body></html>"""


# ── Ejecución de items del checklist + export Excel ─────────────────────

CHECKLIST_DS = "FSCM (Financials/SCM)"


def checklist_query(c, item, env, mode, bu_id=None):
    """Ejecuta la SQL de conteo o detalle de un item en un entorno → (cols, rows)."""
    sql = item["detail_sql"] if mode == "detail" else item["count_sql"]
    base, user, pw = creds(c, env)
    if not base:
        raise RuntimeError(f"Entorno '{env}' sin URL")
    bu_col = bu_column_for(c, CHECKLIST_DS, sql, env) if bu_id else None
    prepared = prepare_sql(sql, user, bu_id, None, bu_col)
    data = fbip.run_sql(base, user, pw, report_for(c, CHECKLIST_DS), prepared)
    return fbip.parse_rows(data)


def checklist_collect(c, items, envs, bu_id=None, workers=8):
    """Conteo + detalle de cada item en cada entorno, en paralelo.
    Devuelve {(índice del item, env): {"count", "count_error", "columns", "rows", "detail_error"}}."""
    from concurrent.futures import ThreadPoolExecutor

    def one(idx, item, env):
        out = {"count": None, "count_error": None, "columns": [], "rows": [], "detail_error": None}
        try:
            cols, rows = checklist_query(c, item, env, "count", bu_id)
            out["count"] = int(rows[0].get(cols[0], 0) or 0) if rows else 0
        except BaseException as e:  # fbip puede lanzar SystemExit
            out["count_error"] = str(e) or e.__class__.__name__
        try:
            out["columns"], out["rows"] = checklist_query(c, item, env, "detail", bu_id)
        except BaseException as e:
            out["detail_error"] = str(e) or e.__class__.__name__
        return (idx, env), out

    with ThreadPoolExecutor(max_workers=workers) as ex:
        # clave por posición: "step" no es único en CHECKLIST (el paso 5 está repetido)
        futs = [ex.submit(one, i, it, env) for i, it in enumerate(items) for env in envs]
        return dict(f.result() for f in futs)


def checklist_build_xlsx(items, envs, data, bu_id=None):
    """Excel con hoja Resumen (conteos por entorno) y hoja Entidades
    (cada registro con su presencia ✓/✗ en cada entorno, misma clave que la vista diff)."""
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3864")
    ok_fill = PatternFill("solid", fgColor="C6EFCE")
    ko_fill = PatternFill("solid", fgColor="FFC7CE")
    warn_fill = PatternFill("solid", fgColor="FFEB9C")

    def header(ws, cols):
        ws.append(cols)
        for cell in ws[1]:
            cell.font, cell.fill = bold, head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"

    def autosize(ws):
        for i, col in enumerate(ws.iter_cols(values_only=True), 1):
            width = max((len(str(v)) for v in col if v is not None), default=8)
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen"
    header(ws, ["Paso", "Módulo", "Tarea", "Responsable", "Método"] + envs + ["Estado"])
    first_env_col = 6
    for i, it in enumerate(items):
        cells = [data[(i, env)] for env in envs]
        counts = [d["count"] for d in cells]
        if any(d["count_error"] for d in cells):
            status = "error"
        elif len(set(counts)) == 1:
            status = "vacío" if counts[0] == 0 else "iguales"
        else:
            status = "diferencias"
        ws.append([it["step"], it["module"], it["task"], it.get("responsible", ""),
                   it.get("config_method") or "UI"]
                  + [d["count"] if not d["count_error"] else "ERROR" for d in cells] + [status])
        r = ws.max_row
        for j, d in enumerate(cells):
            err = d["count_error"] or (d["detail_error"] and "Detalle no disponible: " + d["detail_error"])
            if err:
                ws.cell(r, first_env_col + j).comment = Comment(err[:500], "checklist")
        ws.cell(r, first_env_col + len(envs)).fill = {
            "iguales": ok_fill, "diferencias": warn_fill, "error": ko_fill}.get(status, PatternFill())
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)

    ws = wb.create_sheet("Entidades")
    header(ws, ["Paso", "Módulo", "Tarea", "Campo clave", "Clave", "Contexto"] + envs + ["En todos"])
    first_env_col = 7
    for i, it in enumerate(items):
        cells = [data[(i, env)] for env in envs]
        cols = next((d["columns"] for d in cells if d["columns"]), [])
        if not cols:
            continue
        key_col, ctx_col = cols[0], (cols[1] if len(cols) > 1 else None)
        rows_by_key = {}
        for env, d in zip(envs, cells):
            for row in d["rows"]:
                k = str(row.get(key_col) or "").strip()
                rows_by_key.setdefault(k, {})[env] = row
        for k, per_env in rows_by_key.items():
            sample = next(iter(per_env.values()))
            # None = el entorno falló, no se sabe si el registro existe
            present = [None if d["detail_error"] else env in per_env for env, d in zip(envs, cells)]
            ws.append([it["step"], it["module"], it["task"], key_col, k,
                       sample.get(ctx_col, "") if ctx_col else ""]
                      + ["error" if p is None else "✓" if p else "✗" for p in present]
                      + ["Sí" if all(present) else "No"])
            r = ws.max_row
            for j, p in enumerate(present):
                c = ws.cell(r, first_env_col + j)
                c.fill = warn_fill if p is None else ok_fill if p else ko_fill
                c.alignment = Alignment(horizontal="center")
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)

    if bu_id:
        wb["Resumen"].append([])
        wb["Resumen"].append([f"Filtrado por Business Unit: {bu_id}"])

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ── HTTP handler ───────────────────────────────────────────────────────

class H(BaseHTTPRequestHandler):
    def log_message(self, *a): pass

    def _send(self, code, body, ctype="application/json"):
        raw = body.encode("utf-8") if isinstance(body, str) else body
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", len(raw))
        self.end_headers()
        self.wfile.write(raw)

    def do_GET(self):
        u = urlparse(self.path)
        if u.path in ("/", "/index.html"):
            return self._send(200, PAGE, "text/html; charset=utf-8")

        if u.path == "/api/auth":
            prof = SESSION.get("profile")
            envs = compare_envs()
            missing = [n for n in envs if n not in SESSION["conns"]]
            payload = {"keyring": profiles.keyring_available(), "clients": profiles.list_clients(),
                       "client": None, "logged_in": False}
            if prof:
                env_list = prof.get("environments") or []
                payload.update({
                    "client": {"slug": SESSION["client"], "name": prof.get("name") or SESSION["client"]},
                    "envs": envs, "missing": missing,
                    "logged_in": bool(envs) and not missing,
                    "users": sorted({e.get("user", "") for e in env_list if e.get("name") in envs}),
                    "fbdi_source": prof.get("fbdi_source") or (envs[0] if envs else ""),
                    "fbdi_target": prof.get("fbdi_target") or (envs[-1] if envs else "")})
            return self._send(200, json.dumps(payload))

        if u.path == "/api/clients/get":
            slug = (parse_qs(u.query).get("slug") or [""])[0]
            prof = profiles.load(slug)
            if not prof:
                return self._send(404, json.dumps({"error": f"No existe el cliente '{slug}'."}))
            for e in prof.get("environments") or []:
                e["has_password"] = bool(profiles.get_password(slug, e["name"]))
            prof["slug"] = profiles.slugify(slug)
            return self._send(200, json.dumps(prof))

        if u.path == "/api/checklist/items":
            items = []
            for it in ckl.CHECKLIST:
                spec = fbdi_spec_for_step(it["step"])
                method = it.get("config_method", "")
                has_spec = spec is not None and (("FBDI" in method) or ("REST" in method))
                has_rest = has_spec and bool(fbdi_rest_objects(spec))
                has_bi = has_spec and ("FBDI" in method) and bool(spec.get("sheets"))
                items.append({"module": it["module"], "step": it["step"], "task": it["task"],
                              "responsible": it["responsible"],
                              "config_method": method,
                              "fbdi": has_spec, "fbdi_bi": has_bi, "fbdi_rest": has_rest,
                              "rest_resource": it.get("rest_resource")})
            return self._send(200, json.dumps({"items": items}))

        return self._send(404, json.dumps({"error": "not found"}))

    def do_POST(self):
        u = urlparse(self.path)

        if u.path.startswith(("/api/setup/", "/api/clients/")) or u.path == "/api/logout":
            try:
                n = int(self.headers.get("Content-Length", 0))
                req = json.loads(self.rfile.read(n) or b"{}")
                if u.path == "/api/setup/probe":
                    return self._send(200, json.dumps(profiles.probe(req.get("base"))))
                if u.path == "/api/setup/test":
                    base = profiles.normalize_base(req.get("base"))
                    user, pw = (req.get("user") or "").strip(), req.get("pass") or ""
                    if not pw and req.get("slug") and req.get("env"):  # contraseña ya guardada
                        pw = profiles.get_password(req["slug"], req["env"]) or ""
                    ok, msg = profiles.check_login(base, user, pw)
                    out = {"ok": ok, "message": msg}
                    if ok:
                        out["sql"] = profiles.check_sql(base, user, pw, req.get("sql_report") or None)
                    return self._send(200, json.dumps(out))
                if u.path == "/api/clients/save":
                    prof = req.get("profile") or {}
                    envs = prof.get("environments") or []
                    if not (prof.get("name") or "").strip():
                        return self._send(200, json.dumps({"ok": False, "error": "Falta el nombre del cliente."}))
                    names = [(e.get("name") or "").strip().upper() for e in envs]
                    if not envs or any(not x for x in names) or len(set(names)) != len(names):
                        return self._send(200, json.dumps({"ok": False, "error": "Cada instancia necesita un nombre único."}))
                    for e, nm in zip(envs, names):
                        e["name"], e["base"] = nm, profiles.normalize_base(e.get("base"))
                        if not e["base"]:
                            return self._send(200, json.dumps({"ok": False, "error": f"La instancia {nm} no tiene URL."}))
                    if len([e for e in envs if e.get("compare", True)]) < 2:
                        return self._send(200, json.dumps({"ok": False, "error": "Marca al menos dos instancias para comparar."}))
                    passwords = {(k or "").upper(): v for k, v in (req.get("passwords") or {}).items() if v}
                    remember = bool(req.get("remember")) and profiles.keyring_available()
                    slug = profiles.save(prof, passwords if remember else None)
                    missing = open_client(slug, passwords)
                    return self._send(200, json.dumps({"ok": True, "slug": slug, "missing": missing}))
                if u.path == "/api/clients/open":
                    missing = open_client(req.get("slug") or "", req.get("passwords") or {})
                    return self._send(200, json.dumps({"ok": not missing, "missing": missing}))
                if u.path == "/api/clients/delete":
                    profiles.delete(req.get("slug") or "")
                    if SESSION.get("client") == profiles.slugify(req.get("slug") or ""):
                        SESSION.update(client=None, profile=None, conns={})
                    return self._send(200, json.dumps({"ok": True}))
                if u.path == "/api/logout":
                    SESSION.update(client=None, profile=None, conns={})
                    return self._send(200, json.dumps({"ok": True}))
                return self._send(404, json.dumps({"error": "not found"}))
            except Exception as e:
                return self._send(200, json.dumps({"ok": False, "error": str(e)}))

        if u.path == "/api/checklist/bus":
            try:
                n = int(self.headers.get("Content-Length", 0))
                req = json.loads(self.rfile.read(n) or b"{}")
                env = req.get("env") or (compare_envs() or [""])[0]
                c = cfg()
                ds = "FSCM (Financials/SCM)"
                base, user, pw = creds(c, env)
                if not base:
                    return self._send(200, json.dumps({"bus": [], "error": f"Entorno '{env}' sin URL"}))
                report = report_for(c, ds)
                sql = "SELECT bu_id, bu_name FROM FUN_ALL_BUSINESS_UNITS_V WHERE status='A' ORDER BY bu_name"
                data = fbip.run_sql(base, user, pw, report, sql)
                _, rows = fbip.parse_rows(data)
                bus = [{"id": r.get("BU_ID", ""), "name": r.get("BU_NAME", "")} for r in rows]
                return self._send(200, json.dumps({"bus": bus}))
            except Exception as e:
                return self._send(200, json.dumps({"bus": [], "error": str(e)}))

        if u.path == "/api/checklist/run":
            try:
                n = int(self.headers.get("Content-Length", 0))
                req = json.loads(self.rfile.read(n) or b"{}")
                env = req.get("env")
                step = req.get("step")
                mode = req.get("mode", "count")
                bu_id = req.get("bu")
                item = next((it for it in ckl.CHECKLIST if it["step"] == step), None)
                if not item:
                    return self._send(400, json.dumps({"error": f"Step {step} not found"}))
                cols, rows = checklist_query(cfg(), item, env, mode, bu_id)
                if mode == "count":
                    cnt = rows[0].get(cols[0], "0") if rows else "0"
                    return self._send(200, json.dumps({"count": cnt, "env": env, "step": step}))
                else:
                    return self._send(200, json.dumps({"columns": cols, "rows": rows, "env": env, "step": step}))
            except (RuntimeError, SystemExit) as e:
                return self._send(200, json.dumps({"error": str(e), "count": None}))
            except Exception as e:
                return self._send(200, json.dumps({"error": f"{e}", "count": None}))

        if u.path == "/api/checklist/export":
            try:
                n = int(self.headers.get("Content-Length", 0))
                req = json.loads(self.rfile.read(n) or b"{}")
                bu_id = req.get("bu") or None
                steps = req.get("steps")
                items = [it for it in ckl.CHECKLIST if not steps or it["step"] in steps]
                if not items:
                    return self._send(200, json.dumps({"error": "No hay tareas que exportar."}))
                envs = compare_envs()
                data = checklist_collect(cfg(), items, envs, bu_id)
                raw = checklist_build_xlsx(items, envs, data, bu_id)
                fname = "Fusion_Checklist" + (f"_BU{bu_id}" if bu_id else "") + ".xlsx"
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", f'attachment; filename="{fname}"')
                self.send_header("Content-Length", len(raw))
                self.end_headers()
                self.wfile.write(raw)
                return
            except Exception as e:
                return self._send(200, json.dumps({"error": f"{e}"}))

        if u.path in ("/api/fbdi/delta", "/api/fbdi/download"):
            try:
                n = int(self.headers.get("Content-Length", 0))
                req = json.loads(self.rfile.read(n) or b"{}")
                step = req.get("step")
                source = req.get("source")
                target = req.get("target")
                spec = fbdi_spec_for_step(step)
                if not spec:
                    return self._send(200, json.dumps({"error": f"El paso {step} no tiene plantilla FBDI configurada."}))
                if not source or not target or source == target:
                    return self._send(200, json.dumps({"error": "Origen y destino deben ser entornos distintos."}))
                c = cfg()
                sheets = fbdi_compute_delta(c, spec, source, target)
                if u.path == "/api/fbdi/download":
                    raw = fbdi_build_xlsx(sheets)
                    fname = f"FBDI_{spec['upload_task'].split()[0]}_{source}_to_{target}.xlsx"
                    self.send_response(200)
                    self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                    self.send_header("Content-Disposition", f'attachment; filename="{fname}"')
                    self.send_header("Content-Length", len(raw))
                    self.end_headers()
                    self.wfile.write(raw)
                    return
                payload = {"title": spec["title"], "template_task": spec["template_task"],
                           "upload_task": spec["upload_task"], "note": spec.get("note", ""),
                           "source": source, "target": target,
                           "sheets": [{k: v for k, v in s.items()} for s in sheets]}
                return self._send(200, json.dumps(payload))
            except (RuntimeError, SystemExit) as e:
                return self._send(200, json.dumps({"error": str(e)}))
            except Exception as e:
                return self._send(200, json.dumps({"error": f"{e}"}))

        if u.path in ("/api/fbdi/rest_plan", "/api/fbdi/rest_execute"):
            try:
                n = int(self.headers.get("Content-Length", 0))
                req = json.loads(self.rfile.read(n) or b"{}")
                step = req.get("step")
                source = req.get("source")
                target = req.get("target")
                spec = fbdi_spec_for_step(step)
                if not spec or not fbdi_rest_objects(spec):
                    return self._send(200, json.dumps({"error": f"El paso {step} no tiene carga REST configurada."}))
                if not source or not target or source == target:
                    return self._send(200, json.dumps({"error": "Origen y destino deben ser entornos distintos."}))
                c = cfg()
                if u.path == "/api/fbdi/rest_plan":
                    plan = fbdi_rest_plan(c, spec, source, target)
                    return self._send(200, json.dumps({"source": source, "target": target, "plan": plan}))
                # rest_execute: exige confirmación explícita
                if not req.get("confirm"):
                    return self._send(200, json.dumps({"error": "Falta confirmación (confirm=true)."}))
                result = fbdi_rest_execute(c, spec, source, target)
                result.update({"source": source, "target": target})
                return self._send(200, json.dumps(result))
            except (RuntimeError, SystemExit) as e:
                return self._send(200, json.dumps({"error": str(e)}))
            except Exception as e:
                return self._send(200, json.dumps({"error": f"{e}"}))

        return self._send(404, json.dumps({"error": "not found"}))


# ── main ───────────────────────────────────────────────────────────────

def main():
    srv = ThreadingHTTPServer((HOST, PORT), H)
    url = f"http://{HOST}:{PORT}"
    print(f"Fusion Checklist (multi-cliente) escuchando en {url}")
    if not profiles.keyring_available():
        print("⚠ 'keyring' no disponible: las contraseñas solo se guardarán durante la sesión.")
    try:
        webbrowser.open(url)
    except Exception:
        pass
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        print("\nCerrado.")


if __name__ == "__main__":
    main()
